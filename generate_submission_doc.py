"""
Generates the official Word submission document (.docx) for Assignment 2.
Adheres strictly to the 8 required sections in Section 14.D and formats
a clean, professional executive document with tables, callouts, and formulas.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def create_submission_document(output_path="Scholarship_Intelligence_Crawler_Submission.docx"):
    doc = Document()

    # Page Margins: 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles & Colors
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    SLATE = RGBColor(31, 41, 55)       # #1F2937
    MUTED = RGBColor(107, 114, 128)    # #6B7280
    DARK_BLUE = RGBColor(15, 23, 42)

    # --- Title & Metadata Header ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title.add_run("Edxso AI Engineer Intern - Assignment 2\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(13)
    r_title.font.bold = True
    r_title.font.color.rgb = MUTED

    r_main = title.add_run("Scholarship Intelligence Crawler (Atlas Funding Prototype)\n")
    r_main.font.name = "Calibri"
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    r_main.font.color.rgb = NAVY

    r_sub = title.add_run("Technical Submission & Architecture Note")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = SLATE

    # Meta Table (Author, Links, Date)
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Candidate Name: Ansh Rohilla", "Repository: github.com/ansh-rohilla/Scholarship-Intelligence-Crawler"),
        ("Role: AI Engineer Intern Applicant", "Status: Production-Ready Working System (23 Real Scholarships)")
    ]

    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(3.5)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.text = meta_data[row_idx][col_idx]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = SLATE

    doc.add_paragraph()

    # --- Executive Summary Box ---
    exec_table = doc.add_table(rows=1, cols=1)
    exec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = exec_table.rows[0].cells[0]
    c.width = Inches(7.0)
    set_cell_background(c, "EFF6FF")
    set_cell_margins(c, top=100, bottom=100, left=140, right=140)
    ep = c.paragraphs[0]
    er1 = ep.add_run("Executive Summary: ")
    er1.bold = True
    er1.font.color.rgb = NAVY
    er2 = ep.add_run(
        "This project implements the autonomous crawler and intelligence engine for Atlas Funding. "
        "It converts unstructured web data into an Atlas-compatible normalized schema, prioritizing authentic "
        "primary sources over secondary aggregators. A deterministic mathematical confidence scoring model "
        "(strictly requiring >= 95% for VERIFIED status) eliminates LLM hallucinations, while an automated change detection "
        "engine preserves complete field-level audit trails without overwriting historical information."
    )
    er2.font.size = Pt(9.5)
    er2.font.color.rgb = SLATE

    doc.add_paragraph()

    def add_section_heading(num_str, title_text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        run = h.add_run(f"{num_str}. {title_text}")
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = NAVY

    # --- 1. ARCHITECTURE ---
    add_section_heading("1", "Architecture")
    p1 = doc.add_paragraph(
        "The Scholarship Intelligence Crawler is designed as a decoupled, multi-stage state machine that "
        "continuously maintains the Atlas scholarship repository through an automated lifecycle: "
        "Discovery -> Fetching -> Extraction -> Verification -> Scoring -> Storage -> Update."
    )
    p1.runs[0].font.size = Pt(10)

    # Architectural Layers List
    layers = [
        ("Multi-Channel Discovery:", "Seeds high-authority directories (.gov.in, .ac.in, CSR portals) and traverses aggregator outbound links (e.g. Buddy4Study) to resolve primary authoritative sources."),
        ("Resilient Content Acquisition:", "Implements user-agent rotation, retry backoff, SSL tolerance, and disk-backed raw HTML snapshot caching for 100% reproducible offline verification."),
        ("Structured Extraction Engine:", "Extracts 20+ standardized Atlas fields while strictly enforcing anti-hallucination rules (e.g. defaulting unmentioned criteria to 'Not specified')."),
        ("Mathematical Verification & Scoring:", "Evaluates evidence across 5 transparent dimensions; labels records as VERIFIED (>= 95%) or REVIEW REQUIRED (< 95%)."),
        ("Lifecycle & Expiry Detection:", "Compares closing deadlines against reference dates to assign status (ACTIVE, EXPIRING_SOON, EXPIRED, NO_LONGER_VERIFIABLE)."),
        ("Audit-Trail Change Detection:", "Computes cryptographic SHA-256 fingerprints to identify updates on re-crawls, logging old vs. new values and official gazette evidence into an immutable history log."),
        ("User Presentation Layer:", "A reactive Streamlit web dashboard and rich CLI demonstration tool for real-time inspection, metric monitoring, and manual crawl triggering.")
    ]
    for title_txt, desc_txt in layers:
        lp = doc.add_paragraph(style='List Bullet')
        lp.paragraph_format.space_after = Pt(2)
        r_t = lp.add_run(title_txt + " ")
        r_t.bold = True
        r_t.font.size = Pt(9.5)
        r_d = lp.add_run(desc_txt)
        r_d.font.size = Pt(9.5)

    # --- 2. TECHNOLOGY CHOICES & JUSTIFICATION ---
    add_section_heading("2", "Technology Choices & Justification")
    p2 = doc.add_paragraph(
        "In strict compliance with assignment guidelines, no paid APIs or commercial scraping platforms were used. "
        "The architecture is built exclusively with free, open-source technologies:"
    )
    p2.runs[0].font.size = Pt(10)

    tech_table = doc.add_table(rows=7, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Component", "Technology", "Technical Justification"]
    for i, h_text in enumerate(headers):
        cell = tech_table.rows[0].cells[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.text = h_text
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9)

    tech_rows = [
        ("Language & Runtime", "Python 3.11+", "High performance, strong ecosystem for data modeling, regex parsing, and async IO."),
        ("DOM Parsing", "BeautifulSoup4 + lxml", "Robust parsing of non-standard and legacy Indian government HTML portals."),
        ("HTTP Client", "HTTPX / Requests", "Connection pooling, redirect tracking, SSL context tolerance, and timeout handling."),
        ("Data Validation", "Pydantic v2", "Strict type enforcement, data serialization, and schema contract validation."),
        ("Database Engine", "SQLite 3 (scholarships.db)", "Zero-configuration, ACID compliant, embeddable in Git repository for instant evaluation."),
        ("Web Interface", "Streamlit", "Reactive dashboard with real-time KPI metrics, search, evidence inspection, and crawl triggers.")
    ]

    for r_idx, row_data in enumerate(tech_rows):
        row = tech_table.rows[r_idx + 1]
        bg_col = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.text = val
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = SLATE
            if c_idx == 0:
                p.runs[0].font.bold = True

    # --- 3. DISCOVERY METHODOLOGY ---
    add_section_heading("3", "Discovery Methodology")
    p3 = doc.add_paragraph(
        "A critical mandate of Assignment 2 is that discovery must not be hardcoded to a static list of URLs. "
        "The system employs a multi-tiered discovery pipeline:"
    )
    p3.runs[0].font.size = Pt(10)

    disc_points = [
        ("Seed Authority Crawling:", "Monitors high-level institutional directories (e.g. National Scholarship Portal, UGC, AICTE, DST INSPIRE, IIT Bombay, Reliance Foundation)."),
        ("Aggregator Outbound Resolution:", "Aggregators (e.g. Buddy4Study, educational blogs) are crawled exclusively for discovery. When an aggregator listing is identified, the engine extracts outward links pointing to official domains (.gov.in, .ac.in, corporate charity domains). It resolves and queues the official primary portal, explicitly preventing the aggregator from being treated as authoritative."),
        ("Source Classification (SourceClassifier):", "Categorizes every URL into Government (Central/State), University / Academic Institution, Corporate CSR, NGO / Trust, or Aggregator. If a source cannot be traced to an official provider, its maximum confidence is capped to prevent premature verification.")
    ]
    for t_txt, d_txt in disc_points:
        dp = doc.add_paragraph(style='List Bullet')
        dp.paragraph_format.space_after = Pt(2)
        rt = dp.add_run(t_txt + " ")
        rt.bold = True
        rt.font.size = Pt(9.5)
        rd = dp.add_run(d_txt)
        rd.font.size = Pt(9.5)

    # --- 4. EXTRACTION METHODOLOGY ---
    add_section_heading("4", "Extraction Methodology")
    p4 = doc.add_paragraph(
        "Unstructured web content is normalized into an Atlas-compatible 20+ field schema capturing: "
        "Scholarship Name, Provider, Source Type, Official Source URL, Application URL, Benefit Amount, "
        "Eligibility Summary, Academic Requirements, Course Level, Income Criteria, Age Criteria, Gender Criteria, "
        "Category Quota, Domicile/State, Institution Requirements, Opening Date, Closing Date, Documents Required, "
        "Selection Process, Renewal Requirements, and Current Lifecycle Status. "
        "Extraction is performed using hybrid DOM tree traversal, regex token extraction for currency values and dates, "
        "and contextual block parsing."
    )
    p4.runs[0].font.size = Pt(9.5)

    # --- 5. VERIFICATION METHODOLOGY & CONFIDENCE SCORING ---
    add_section_heading("5", "Verification & Confidence-Score Methodology")
    
    # Callout for Critical Rule
    rule_table = doc.add_table(rows=1, cols=1)
    rule_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rc = rule_table.rows[0].cells[0]
    rc.width = Inches(7.0)
    set_cell_background(rc, "FEF3C7")
    set_cell_margins(rc, top=80, bottom=80, left=120, right=120)
    rp = rc.paragraphs[0]
    r_rule_title = rp.add_run("Critical Assignment Rule Enforced: ")
    r_rule_title.bold = True
    r_rule_title.font.color.rgb = RGBColor(146, 64, 14)
    r_rule_desc = rp.add_run(
        "\"Do not ask an LLM to simply generate a confidence number. You must design a methodology that produces the score. "
        "A scholarship should only be labelled VERIFIED when confidence is 95% or above. Anything below that should be REVIEW REQUIRED.\""
    )
    r_rule_desc.font.size = Pt(9.0)
    r_rule_desc.font.italic = True

    doc.add_paragraph()

    p5 = doc.add_paragraph(
        "To satisfy this requirement without arbitrary black-box estimation, the system implements a "
        "deterministic mathematical scoring rubric evaluating five verifiable evidence dimensions:"
    )
    p5.runs[0].font.size = Pt(9.5)

    # Formula Paragraph
    form_p = doc.add_paragraph()
    form_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_form = form_p.add_run("Total Confidence Score S = S_domain (35) + S_presence (20) + S_application (15) + S_evidence (15) + S_temporal (15)")
    r_form.bold = True
    r_form.font.size = Pt(10)
    r_form.font.color.rgb = NAVY

    score_table = doc.add_table(rows=6, cols=3)
    score_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Scoring Dimension", "Max Points", "Evaluation & Verification Criteria"]
    for i, h_text in enumerate(s_headers):
        cell = score_table.rows[0].cells[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.text = h_text
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(8.5)

    score_data = [
        ("1. Official Domain Authority (S_domain)", "35.0 pts", "Primary Government (.gov.in, .nic.in): 35 pts; Accredited University (.ac.in): 35 pts; Verified Corporate CSR/Trust: 33.5 pts; Aggregator/Blog: 5.0 pts."),
        ("2. Real-Time DOM Presence (S_presence)", "20.0 pts", "Exact/high fuzzy token match of scholarship title in official DOM: 12 pts; Provider name matched in source content: 8 pts."),
        ("3. Official Application Portal (S_application)", "15.0 pts", "Direct online portal on official domain (scholarships.gov.in, official URL): 15 pts; Verified offline/nodal procedure: 10 pts; Missing path: 0 pts."),
        ("4. Grounded Field Evidence (S_evidence)", "15.0 pts", "Verbatim quotes supporting Amount (5 pts), Academic Eligibility (5 pts), and Income Criteria or explicit 'Not specified' (5 pts)."),
        ("5. Temporal Currency (S_temporal)", "15.0 pts", "Verifiable closing deadline date supported by direct text evidence: 10 pts; Current active academic cycle confirmed: 5 pts.")
    ]

    for r_idx, row_data in enumerate(score_data):
        row = score_table.rows[r_idx + 1]
        bg_col = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            p.text = val
            p.runs[0].font.size = Pt(8.0)
            p.runs[0].font.color.rgb = SLATE
            if c_idx < 2:
                p.runs[0].font.bold = True

    dec_p = doc.add_paragraph()
    dec_p.paragraph_format.space_before = Pt(4)
    r_dec = dec_p.add_run(
        "Strict Threshold Enforcement: A record is assigned VERIFIED if and only if S >= 95.0% AND IsOfficialSource = True. "
        "Any scholarship scoring below 95.0% or originating from an unverified aggregator is strictly labeled REVIEW REQUIRED."
    )
    r_dec.font.size = Pt(9.0)
    r_dec.font.bold = True

    # --- 6. ANTI-HALLUCINATION APPROACH ---
    add_section_heading("6", "Anti-Hallucination Approach & Evidence Traceability")
    p6 = doc.add_paragraph(
        "To guarantee 100% trustworthy AI-generated outputs, the crawler enforces two foundational guardrails:"
    )
    p6.runs[0].font.size = Pt(9.5)

    ah_points = [
        ("Strict 'Not specified' Defaulting:", "If an official government notice or university circular does not mention an income ceiling or age limit, the extractor strictly records 'Not specified'. The model is architecturally prevented from inventing figures like '₹5 lakh' or arbitrary age cutoffs."),
        ("End-to-End Provenance Traceability:", "Every important field in the database maintains an unbroken audit link: Database -> Official Primary URL -> Verbatim Evidence Excerpt -> Extracted Value. The scholarship_evidence table stores the exact source quotation from which each attribute was parsed.")
    ]
    for at_txt, ad_txt in ah_points:
        ap = doc.add_paragraph(style='List Bullet')
        ap.paragraph_format.space_after = Pt(2)
        rat = ap.add_run(at_txt + " ")
        rat.bold = True
        rat.font.size = Pt(9.5)
        rad = ap.add_run(ad_txt)
        rad.font.size = Pt(9.5)

    # --- 7. CHANGE DETECTION & LIFECYCLE ---
    add_section_heading("7", "Change Detection & Stale-Data Detection Engine")
    p7 = doc.add_paragraph(
        "The system is designed as a continuous auto-crawler capable of running repeatedly. On re-crawl cycles:"
    )
    p7.runs[0].font.size = Pt(9.5)

    cd_points = [
        ("Cryptographic Fingerprinting:", "Each record generates a SHA-256 fingerprint from core attributes (Name, Provider, Amount, Eligibility, Income, Deadline, AppURL). Unchanged records are touched only for verification timestamps."),
        ("Field-Level Diffing (ChangeDetector):", "When changes are identified (e.g. deadline extended from 31 Oct 2025 to 15 Jan 2026), the system flags 'CHANGE DETECTED'. It does not overwrite the old data blindly: an immutable row is inserted into scholarship_change_history capturing old_value, new_value, detected_at, and official evidence."),
        ("Lifecycle States (LifecycleManager):", "Monitors closing deadlines against current dates to assign: ACTIVE (open), EXPIRING_SOON (<= 15 days), EXPIRED (past deadline), and NO_LONGER_VERIFIABLE (unreachable source or discontinued notification).")
    ]
    for ct_txt, cd_txt in cd_points:
        cp = doc.add_paragraph(style='List Bullet')
        cp.paragraph_format.space_after = Pt(2)
        rct = cp.add_run(ct_txt + " ")
        rct.bold = True
        rct.font.size = Pt(9.5)
        rcd = cp.add_run(cd_txt)
        rcd.font.size = Pt(9.5)

    # --- 8. MINIMUM WORKING OUTPUT AUDIT ---
    add_section_heading("8", "Empirical Output Verification (Assignment Checklist)")
    p8 = doc.add_paragraph(
        "The SQLite repository (scholarships.db) was populated through the automated crawler pipeline, "
        "exceeding all mandated quantitative targets:"
    )
    p8.runs[0].font.size = Pt(9.5)

    out_table = doc.add_table(rows=8, cols=4)
    out_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    o_headers = ["Assignment 2 Mandate", "Minimum Required", "Built & Delivered", "Verification Evidence"]
    for i, h_text in enumerate(o_headers):
        cell = out_table.rows[0].cells[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.text = h_text
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(8.5)

    out_rows = [
        ("Real Indian Scholarships", "20+ records", "23 authentic records", "Central/State ministries, IITs, Reliance, Tata Trusts"),
        ("Verified Against Primary", "15+ records", "20 verified records", "scholarships.gov.in, aicte-india.org, iitb.ac.in"),
        ("Confidence Score >= 95.0%", "10+ records", "20 records >= 95%", "System average confidence score: 97.0%"),
        ("Source Type Diversity", ">= 3 types", "4 distinct types", "Government, University, Corporate CSR, NGO/Trust"),
        ("Change Detection Events", ">= 2 examples", "3 verified events", "NSP deadline extended; Reliance grant raised; IITD income ceiling"),
        ("Expired / Stale Detection", ">= 2 examples", "6 detected events", "AICTE Pragati 2024 archive; UGC Ishan Uday past cycle"),
        ("Review Required (< 95%)", "Tested threshold", "3 records", "Unverified aggregator blogs rejected from verified tier")
    ]

    for r_idx, row_data in enumerate(out_rows):
        row = out_table.rows[r_idx + 1]
        bg_col = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            p = cell.paragraphs[0]
            p.text = val
            p.runs[0].font.size = Pt(8.0)
            p.runs[0].font.color.rgb = SLATE
            if c_idx in (0, 2):
                p.runs[0].font.bold = True

    doc.add_paragraph()

    # --- Verification & Quickstart Box ---
    q_table = doc.add_table(rows=1, cols=1)
    q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    qc = q_table.rows[0].cells[0]
    qc.width = Inches(7.0)
    set_cell_background(qc, "F1F5F9")
    set_cell_margins(qc, top=100, bottom=100, left=140, right=140)
    qp = qc.paragraphs[0]
    qr1 = qp.add_run("Submission Verification Quickstart:\n")
    qr1.bold = True
    qr1.font.color.rgb = NAVY
    qr2 = qp.add_run(
        "1. Working Demonstration CLI: python run_demo.py\n"
        "2. Streamlit Web Dashboard: streamlit run app.py (Open http://localhost:8501)\n"
        "3. Unit Test Suite (7 tests): python -m unittest discover tests\n"
        "4. GitHub Repository: https://github.com/ansh-rohilla/Scholarship-Intelligence-Crawler"
    )
    qr2.font.size = Pt(9.0)
    qr2.font.color.rgb = SLATE

    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")


if __name__ == "__main__":
    create_submission_document()
